import textwrap

import pytest
from django.core.management import call_command
from docx import Document

from apps.locations.models import City, Country, Town
from apps.seeding.parsers.locations_parser import parse_and_load
from apps.seeding.parsers.price_list_parser import parse

pytestmark = pytest.mark.django_db


def _write_csv(tmp_path, rows):
    path = tmp_path / "sample_cities_and_towns.csv"
    path.write_text("Country,City,Town\n" + "\n".join(rows) + "\n")
    return path


def _write_docx(tmp_path, name, header):
    path = tmp_path / name
    document = Document()
    document.add_paragraph(header)
    document.add_paragraph("Both displayed Minimum Price and Maximum Price")
    document.add_paragraph("Any")
    document.save(path)
    return path


def test_locations_parser_creates_hierarchy(tmp_path):
    csv_path = _write_csv(tmp_path, [
        "Kenya,Nairobi,Westlands",
        "Kenya,Nairobi,Karen",
        "Kenya,Mombasa,Nyali",
    ])

    counts = parse_and_load(csv_path)

    assert counts == {"countries": 1, "cities": 2, "towns": 3}
    assert Country.objects.count() == 1
    assert City.objects.count() == 2
    assert Town.objects.count() == 3
    assert Country.objects.get().currency_code == "KES"


def test_locations_parser_is_idempotent(tmp_path):
    csv_path = _write_csv(tmp_path, ["Kenya,Nairobi,Westlands"])

    parse_and_load(csv_path)
    counts = parse_and_load(csv_path)

    assert counts == {"countries": 0, "cities": 0, "towns": 0}
    assert Town.objects.count() == 1


def test_price_list_parser_reads_header_only(tmp_path):
    path = _write_docx(
        tmp_path, "kenya_sale_price_list.docx",
        "Sale Price List (500,000 KSh to 1,000,000,000 KSh in increments of 100,000)",
    )

    result = parse(path)

    assert result == {
        "country": "Kenya", "listing_type": "sale",
        "currency_code": "KES", "min_price": 500_000, "max_price": 1_000_000_000,
    }


def test_price_list_parser_maps_currency_symbols(tmp_path):
    path = _write_docx(
        tmp_path, "uganda_rent_price_list.docx",
        "Uganda Rent Price List (300,000 UGX to 30,000,000 UGX in increments of 150,000)",
    )

    result = parse(path)

    assert result["currency_code"] == "UGX"
    assert result["listing_type"] == "rent"


def test_seed_locations_command_is_idempotent(monkeypatch, tmp_path):
    kenya_dir = tmp_path / "kenya"
    kenya_dir.mkdir()
    (kenya_dir / "kenya_cities_and_towns.csv").write_text(
        "Country,City,Town\nKenya,Nairobi,Westlands\n"
    )

    import apps.seeding.management.commands.seed_locations as cmd

    monkeypatch.setattr(cmd, "DATA_SOURCE_DIR", tmp_path)
    monkeypatch.setattr(cmd, "COUNTRY_ORDER", ["kenya"])

    call_command("seed_locations")
    call_command("seed_locations")

    assert Country.objects.count() == 1
    assert Town.objects.count() == 1
